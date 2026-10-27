#!/usr/bin/env python3
"""Build a Power BI web-upload workbook (one Excel Table per sheet) from the DuckDB core/mart layer.
SYNTHETIC SAMPLE DATA ONLY - not real Optima / Emeria figures.
Default (no arguments): reads db/berlin_property_ops.duckdb + qa/dq_report.md and writes
powerbi/Berlin_Property_Ops_33k.xlsx (the file property.pbix imports).
usage: build_xlsx.py [<duckdb> <dq_report.md> <out.xlsx> <label>]"""
import os, sys, re, warnings, datetime as dt, decimal, duckdb, openpyxl
warnings.filterwarnings("ignore", message="In write-only mode you must add table columns manually")  # columns are added below
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
from openpyxl.styles import Font
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
os.chdir(ROOT)
if len(sys.argv) >= 5:
    DB, DQMD, OUT, LABEL = sys.argv[1:5]
else:
    DB, DQMD = os.path.join("db", "berlin_property_ops.duckdb"), os.path.join("qa", "dq_report.md")
    OUT, LABEL = os.path.join("powerbi", "Berlin_Property_Ops_33k.xlsx"), "33,000 units / 680 properties"
os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
EXCEL_MAX = 1_048_576 - 1          # data rows per sheet (minus header)
con = duckdb.connect(DB, read_only=True)

# ---- payments folded onto rent lines (fact_payment itself is replaced, see WEB_BUILD_CHECKLIST.md)
con.execute("""
CREATE TEMP TABLE rr AS
SELECT r.period_month, r.due_date, r.unit_id, r.tenant_id,
       r.kaltmiete_due, r.nk_prepayment_due, r.total_due,
       COALESCE(SUM(CASE WHEN p.payment_date <= r.due_date THEN p.amount END), 0)              AS paid_by_due_date,
       COALESCE(SUM(CASE WHEN p.payment_date <= last_day(r.period_month) THEN p.amount END), 0) AS paid_by_month_end,
       COALESCE(SUM(p.amount), 0)                                                              AS paid_to_snapshot
FROM core.fact_rent_roll r LEFT JOIN core.fact_payment p USING (rent_roll_id)
WHERE r.period_month >= DATE '2024-01-01'
GROUP BY ALL ORDER BY r.period_month, r.unit_id""")
con.execute("""
CREATE TEMP TABLE pm AS
SELECT r.property_id, CAST(date_trunc('month', p.payment_date) AS DATE) AS payment_month, p.payment_method,
       COUNT(*) AS payments, SUM(p.amount) AS amount
FROM core.fact_payment p JOIN core.fact_rent_roll r USING (rent_roll_id)
GROUP BY ALL ORDER BY 2, 1, 3""")
n_rr = con.sql("SELECT COUNT(*) FROM rr").fetchone()[0]

tables = [(t, f"SELECT * FROM core.{t}") for t in
          ["dim_property", "dim_unit", "dim_tenant", "dim_date", "dim_contractor", "dim_priority"]]
if n_rr <= EXCEL_MAX:
    tables.append(("fact_rent_roll", "SELECT * FROM rr"))
else:  # split by year; the web model re-unions them into fact_rent_roll (calculated table)
    tables.append(("fact_rent_roll_2024", "SELECT * FROM rr WHERE year(period_month) = 2024"))
    tables.append(("fact_rent_roll_2025_2026", "SELECT * FROM rr WHERE year(period_month) >= 2025"))
tables += [("fact_payment_monthly", "SELECT * FROM pm")]
tables += [(t, f"SELECT * FROM core.{t}") for t in
           ["fact_maintenance_ticket", "fact_vacancy_reletting", "fact_service_charge_statement",
            "fact_service_charge_cost", "fact_energy_monthly"]]
tables += [("mart_rent_arrears_monthly", "SELECT * FROM mart.rent_arrears_monthly ORDER BY month_start, property_id"),
           ("mart_occupancy_monthly", "SELECT * FROM mart.occupancy_monthly ORDER BY month_start, property_id")]

# ---- DQ results parsed from qa/dq_report.md
DESC = {"DQ01": "ticket_id must be unique in the maintenance ticket extract",
        "DQ02": "resolved_at must not be earlier than created_at",
        "DQ03": "every unit needs a living/usable area > 0 m2",
        "DQ04": "every rent-roll unit_id must exist in the unit master",
        "DQ05": "a bank reference may be booked only once",
        "DQ06": "a service-charge statement cannot be issued before its billing period ends",
        "DQ07": "statement header total = sum of its cost lines (tolerance EUR 1)",
        "DQ08": "every payment must reference an existing rent line",
        "DQ09": "total_due = Kaltmiete + NK prepayment",
        "DQ10": "status 'Erledigt' needs resolved_at; no resolved_at after the snapshot",
        "DQ11": "no overlapping tenancies on the same unit",
        "DQ12": "tickets must reference a known property/unit"}
md = open(DQMD, encoding="utf-8").read()
dq_rows, rc_rows = [], []
sec = None
for line in md.splitlines():
    if line.startswith("## "):
        sec = line
    if not line.startswith("| ") or line.startswith("|---") or "check_id" in line or "raw_rows" in line or "injected keys" in line:
        continue
    c = [x.strip() for x in line.strip("|").split("|")]
    if sec and sec.startswith("## Check results"):
        dq_rows.append([c[0], c[1], DESC[c[0]], c[2], c[3], int(c[4]), int(c[5]), c[6], c[7]])
    elif sec and sec.startswith("## Rows before"):
        rc_rows.append([c[0], int(c[1]), int(c[2]), int(c[1]) - int(c[2])])
assert len(dq_rows) == 12, dq_rows

def conv(v):
    if isinstance(v, decimal.Decimal):
        return float(v)
    return v

wb = openpyxl.Workbook(write_only=True)
style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
def add_sheet(name, header, rows_iter):
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    n = 0
    for r in rows_iter:
        ws.append([conv(v) for v in r]); n += 1
    assert n <= EXCEL_MAX, (name, n)
    last = openpyxl.utils.get_column_letter(len(header))
    t = Table(displayName=name, ref=f"A1:{last}{max(n, 1) + 1}", tableStyleInfo=style)
    t.tableColumns = [TableColumn(id=i + 1, name=h) for i, h in enumerate(header)]
    ws.add_table(t)
    return n

readme = wb.create_sheet("README")
lines = [
    "Berlin Property Operations Analytics - SYNTHETIC SAMPLE DATA",
    "",
    f"Scale: {LABEL}.",
    "ALL DATA IN THIS WORKBOOK IS SYNTHETIC (generated by a seeded Python script, seed 20261017).",
    "It is NOT real data from Hausverwaltung Optima, Emeria or any other company; no real tenants, properties or contractors.",
    "Snapshot date of the sample: 2026-09-30. Reporting window 2024-01 .. 2026-09.",
    "",
    "Each sheet holds one Excel Table named exactly like the sheet; Power BI imports the Tables.",
    "fact_payment is not included at payment level: payments are folded onto each rent line",
    "(paid_by_due_date, paid_by_month_end, paid_to_snapshot) and a monthly cash table fact_payment_monthly is provided.",
    "fact_rent_roll holds rent lines from 2024-01 only (2023 warm-up lines are reflected in the arrears mart).",
    "If fact_rent_roll is split into _2024 and _2025_2026 sheets, re-union them in the model (see WEB_BUILD_CHECKLIST.md).",
    "dq_results / dq_row_counts come from qa/dq_report.md (raw-layer checks; the other tables are the cleaned core layer).",
    "_Measures is an empty helper table (one dummy column) that holds the DAX measures.",
]
for l in lines:
    readme.append([l])

counts = {}
for name, q in tables:
    rel = con.sql(q)
    header = rel.columns
    cur = con.execute(q)
    def it():
        while True:
            chunk = cur.fetchmany(50000)
            if not chunk:
                return
            yield from chunk
    counts[name] = add_sheet(name, header, it())
    print(f"{name:32s} {counts[name]:>9,}", flush=True)
counts["dq_results"] = add_sheet("dq_results", ["check_id", "check_name", "description", "table_name", "severity",
                                                "failed_keys", "rows_affected", "result", "handling"], dq_rows)
counts["dq_row_counts"] = add_sheet("dq_row_counts", ["table_name", "raw_rows", "core_rows", "rows_removed"], rc_rows)
counts["_Measures"] = add_sheet("_Measures", ["Dummy"], [[None]])
wb.save(OUT)
print("saved", OUT)
